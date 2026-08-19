#!/usr/bin/env python3
"""
Loads DB_SP_INF.xlsx / "DB_SP_INF  II.xlsx" into DB_SP_INF, and
TB_M_USER_PROC_ID.xlsx into DB_SP_SP.dbo.TB_M_USER_PROC_ID.

Usage:
    pip install -r requirements.txt
    python3 load_excel_data.py           # dry run, no writes
    python3 load_excel_data.py --commit  # actually insert

Env vars (all optional, defaults match docker-compose.yml):
    MSSQL_HOST, MSSQL_PORT, MSSQL_SA_PASSWORD
"""
import os
import sys
import datetime
from pathlib import Path

import openpyxl
import pymssql

REPO_ROOT = Path(__file__).resolve().parent.parent

HOST = os.environ.get("MSSQL_HOST", "localhost")
PORT = int(os.environ.get("MSSQL_PORT", "1433"))
SA_PASSWORD = os.environ.get("MSSQL_SA_PASSWORD", "YourStrong!Passw0rd")

CREATED_BY = "1"
UPDATED_BY = "1"

EXCEL_EPOCH = datetime.datetime(1899, 12, 30)
DRY_RUN = "--commit" not in sys.argv

DATETIME_TYPES = {"datetime", "datetime2", "smalldatetime", "date", "datetimeoffset"}
TEXTUAL_TYPES = {"char", "varchar", "nchar", "nvarchar", "text", "ntext"}

INF_FILES = {
    "DB_SP_INF.xlsx": [
        "TB_I_H01_06_SUPP", "TB_I_B03_PO_FILE_DOMSUPP", "TB_I_B04_PO_FILE_DIRCVND",
        "TB_I_H01_03_PROC_BASIC", "TB_I_H01_02_PNO", "TB_I_H01_12_STORAGE",
        "TB_I_H01_24_PROC_DELV_CTRL", "TB_I_H01_40_KANBAN", "TB_I_H01_08_COST",
        "TB_I_H01_07_PICK_OPS_PNO", "TB_I_H01_34_TOP6_KATASHIKI", "TB_I_H02_58_DDD_PO_INFO",
        "TB_I_H02_59_DDD_INFO_SUM", "TB_I_H02_60_DDD_INFO_DTL", "TB_I_H02_64_BO_ETD_FILE",
        "TB_I_H01_21_CUSTOMER", "TB_I_H01_25_DESTINATION", "TB_I_H01_09_PROC_CTL_PO_CTL",
        "TB_I_H02_35_PNO_STK_DISTR", "TB_I_H02_01_STOCK_FILE", "TB_I_B12_ACCOUNT_PAYABLE",
    ],
    "DB_SP_INF  II.xlsx": [
        "TB_I_H02_04_RECV_UPD_FILE", "TB_I_H02_04_RECV_UPD_FILE_N",
        "TB_I_H02_04_RECV_UPD_FILE_CMPLT", "TB_I_H02_08_RCV_OPS_DTL", "TB_I_H02_08_RCV_OPS_DTL_M",
    ],
}
USER_PROC_ID_FILE = "TB_M_USER_PROC_ID.xlsx"


def get_schema(conn, db, table):
    cur = conn.cursor(as_dict=True)
    cur.execute(f"""
        SELECT c.name, t.name AS type, c.is_nullable
        FROM sys.columns c
        JOIN sys.types t ON c.user_type_id = t.user_type_id
        WHERE c.object_id = OBJECT_ID('{db}.dbo.{table}')
        ORDER BY c.column_id
    """)
    return {r["name"]: (r["type"], bool(r["is_nullable"])) for r in cur.fetchall()}


def read_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # title row
    header = [str(h) for h in next(rows) if h is not None]
    data = []
    for r in rows:
        if all(v is None for v in r):
            continue
        data.append(r[: len(header)])
    return header, data


def convert_value(value, col_type, nullable):
    if value is None:
        if not nullable and col_type in TEXTUAL_TYPES:
            return ""
        return None
    if col_type in DATETIME_TYPES and isinstance(value, (int, float)):
        return EXCEL_EPOCH + datetime.timedelta(days=float(value))
    if col_type in TEXTUAL_TYPES and not isinstance(value, str):
        return str(value)
    return value


def load_inf_table(conn, table, file_path, stats):
    header, data = read_sheet(file_path, table)
    schema = get_schema(conn, "DB_SP_INF", table)

    missing = [h for h in header if h not in schema]
    if missing:
        print(f"  ERROR {table}: columns in excel not found in DB: {missing}")
        stats["errors"] += 1
        return

    col_list = ", ".join(f"[{h}]" for h in header)
    placeholders = ", ".join(["%s"] * len(header))
    sql = f"INSERT INTO dbo.[{table}] ({col_list}) VALUES ({placeholders})"

    rows_out = []
    null_violations = 0
    for row in data:
        converted = []
        for h, v in zip(header, row):
            col_type, nullable = schema[h]
            cv = convert_value(v, col_type, nullable)
            if cv is None and not nullable:
                null_violations += 1
            converted.append(cv)
        rows_out.append(tuple(converted))

    print(f"  {table}: {len(rows_out)} rows, null_violations={null_violations}")
    stats["rows"] += len(rows_out)
    if null_violations:
        stats["errors"] += 1
        return

    if not DRY_RUN:
        cur = conn.cursor()
        cur.executemany(sql, rows_out)
        conn.commit()
        print(f"    -> inserted {len(rows_out)} rows into DB_SP_INF.dbo.{table}")


def load_user_proc_id(conn):
    header, data = read_sheet(REPO_ROOT / USER_PROC_ID_FILE, "Sheet1")
    print(f"TB_M_USER_PROC_ID: header={header}, rows={len(data)}")

    now = datetime.datetime.now(datetime.timezone.utc)
    col_list = "ID, USER_ID_WINDOWS, PROC_CODE_ID, USER_NAME, DEPT_SECTION, CREATED_BY, UPDATED_BY, CREATED_DATE, UPDATED_DATE"
    placeholders = ", ".join(["%s"] * 9)
    sql = f"INSERT INTO dbo.TB_M_USER_PROC_ID ({col_list}) VALUES ({placeholders})"

    rows_out = []
    for row in data:
        rid, user_id_windows, proc_code_id, user_name, dept_section = row[:5]
        rows_out.append((
            rid, user_id_windows, proc_code_id, user_name, dept_section,
            CREATED_BY, UPDATED_BY, now, now,
        ))

    print(f"  {len(rows_out)} rows ready (ID {min(r[0] for r in rows_out)}-{max(r[0] for r in rows_out)})")

    if not DRY_RUN:
        cur = conn.cursor()
        cur.execute("SET IDENTITY_INSERT dbo.TB_M_USER_PROC_ID ON")
        cur.executemany(sql, rows_out)
        cur.execute("SET IDENTITY_INSERT dbo.TB_M_USER_PROC_ID OFF")
        conn.commit()
        print(f"    -> inserted {len(rows_out)} rows into DB_SP_SP.dbo.TB_M_USER_PROC_ID")


def main():
    print(f"=== {'DRY RUN' if DRY_RUN else 'COMMIT'} ({HOST}:{PORT}) ===")

    stats = {"rows": 0, "errors": 0}

    print("--- DB_SP_INF tables ---")
    conn_inf = pymssql.connect(server=HOST, port=PORT, user="sa", password=SA_PASSWORD, database="DB_SP_INF", autocommit=False)
    for file_name, tables in INF_FILES.items():
        file_path = REPO_ROOT / file_name
        for table in tables:
            load_inf_table(conn_inf, table, file_path, stats)
    conn_inf.close()

    print("--- DB_SP_SP.dbo.TB_M_USER_PROC_ID ---")
    conn_sp = pymssql.connect(server=HOST, port=PORT, user="sa", password=SA_PASSWORD, database="DB_SP_SP", autocommit=False)
    load_user_proc_id(conn_sp)
    conn_sp.close()

    print(f"=== done: total_rows={stats['rows']} errors={stats['errors']} ===")
    if stats["errors"] and not DRY_RUN:
        sys.exit(1)


if __name__ == "__main__":
    main()
