#!/usr/bin/env python3
"""
Loads the 20260909 DB_SP_INF workbook snapshot into DB_SP_INF and DB_SP_WRH.

Sheet name == target table name in every case. Each job carries a mode:
  "replace" — the sheet is a full snapshot; empty the table first.
  "append"  — the sheet is a daily delta covering a date window the table does
              not hold yet; insert on top of what is there.

Usage:
    .venv/bin/python3 scripts/load_excel_20260909.py            # dry run, no writes
    .venv/bin/python3 scripts/load_excel_20260909.py --commit   # truncate + insert

Connection comes from an env file (default ../.env, override with --env-file),
falling back to real env vars, falling back to defaults matching
docker-compose.yml:

    MSSQL_HOST=localhost          (or DB_HOST)
    MSSQL_PORT=1433               (or DB_PORT)
    MSSQL_USER=sa                 (or DB_USER)
    MSSQL_SA_PASSWORD=...         (or MSSQL_PASSWORD / DB_PASSWORD)

Never pass the password on the command line — it lands in shell history and in
the process list. Put it in an env file; .env* is gitignored.
"""
import argparse
import datetime
import os
import sys
from pathlib import Path

import openpyxl
import pymssql
from dotenv import load_dotenv

REPO_ROOT = Path(__file__).resolve().parent.parent

# populated by main() once --env-file is known
HOST = PORT = USER = PASSWORD = None

DEFAULT_WORKBOOK = Path.home() / "Downloads" / "DB_SP_INF 20260909.xlsx"

# (db_name, sheet/table name, mode) — sheet name equals the target table name.
JOBS = [
    ("DB_SP_INF", "TB_I_H02_58_DDD_PO_INFO", "replace"),
    ("DB_SP_INF", "TB_I_H02_59_DDD_INFO_SUM", "replace"),
    ("DB_SP_INF", "TB_I_H02_60_DDD_INFO_DTL", "replace"),
    ("DB_SP_INF", "TB_I_H02_49_POOLED_KANBAN", "replace"),
    ("DB_SP_INF", "TB_I_B16_KANBAN", "replace"),
    ("DB_SP_INF", "TB_I_B17_SUPP_DN", "replace"),
    # B03/B04 hold PO_DATE 20260730..20260807 from the Aug 18 workbook; the
    # 20260909 sheets are a single-day 20260909 extract with zero key overlap.
    ("DB_SP_INF", "TB_I_B03_PO_FILE_DOMSUPP", "append"),
    ("DB_SP_INF", "TB_I_B04_PO_FILE_DIRCVND", "append"),
    ("DB_SP_WRH", "TB_M_DEPOT", "replace"),
    ("DB_SP_WRH", "TB_M_SUPPLIER", "replace"),
    ("DB_SP_WRH", "TB_M_KANBAN_SUPPLIER", "replace"),
]

EXCEL_EPOCH = datetime.datetime(1899, 12, 30)

# Rows per INSERT statement. pymssql's executemany issues one round trip per
# row, which dominates runtime on a high-latency link (~19 rows/s to Azure SEA).
# Batching row constructors into one statement cuts round trips ~100x. Capped
# by SQL Server's 1000-row VALUES limit and by keeping the statement text sane.
VALUES_LIMIT = 1000
STMT_VALUE_BUDGET = 8000


def rows_per_statement(ncols):
    return max(1, min(VALUES_LIMIT, STMT_VALUE_BUDGET // max(1, ncols)))

DATETIME_TYPES = {"datetime", "datetime2", "smalldatetime", "date", "datetimeoffset"}
TEXTUAL_TYPES = {"char", "varchar", "nchar", "nvarchar", "text", "ntext"}
NUMERIC_TYPES = {"numeric", "decimal", "int", "bigint", "smallint", "tinyint", "float", "real", "money"}


def load_connection(env_file):
    """Reads connection settings from env_file, then real env vars, then defaults."""
    global HOST, PORT, USER, PASSWORD
    path = Path(env_file)
    if not path.is_file():
        sys.exit(f"env file not found: {path}")
    load_dotenv(path, override=True)

    def pick(*names, default=None):
        for n in names:
            v = os.environ.get(n)
            if v not in (None, ""):
                return v
        return default

    HOST = pick("MSSQL_HOST", "DB_HOST", default="localhost")
    PORT = int(pick("MSSQL_PORT", "DB_PORT", default="1433"))
    USER = pick("MSSQL_USER", "DB_USER", default="sa")
    PASSWORD = pick("MSSQL_SA_PASSWORD", "MSSQL_PASSWORD", "DB_PASSWORD")
    if not PASSWORD:
        sys.exit(f"no password (MSSQL_SA_PASSWORD / MSSQL_PASSWORD / DB_PASSWORD) in {path}")


def connect(db):
    return pymssql.connect(server=HOST, port=PORT, user=USER, password=PASSWORD,
                           database=db, autocommit=False)


def get_schema(conn, db, table):
    cur = conn.cursor(as_dict=True)
    cur.execute(f"""
        SELECT c.name, t.name AS type, c.max_length, c.is_nullable, c.is_identity
        FROM sys.columns c
        JOIN sys.types t ON c.user_type_id = t.user_type_id
        WHERE c.object_id = OBJECT_ID('{db}.dbo.{table}')
        ORDER BY c.column_id
    """)
    rows = cur.fetchall()
    return {
        r["name"]: {
            "type": r["type"],
            # max_length is bytes; nchar/nvarchar store 2 bytes per char
            "chars": (r["max_length"] // 2 if r["type"].startswith("n") and r["max_length"] > 0
                      else r["max_length"]),
            "nullable": bool(r["is_nullable"]),
            "identity": bool(r["is_identity"]),
        }
        for r in rows
    }


def read_sheet(wb, sheet_name):
    """Returns (header, rows). Tolerates sheets with or without a leading title row."""
    ws = wb[sheet_name]
    header = None
    data = []
    for row in ws.iter_rows(values_only=True):
        non_null = [v for v in row if v is not None]
        if header is None:
            # a title row holds a single cell (the table name); the header holds many
            if len(non_null) > 1:
                header = [str(h).strip() for h in row if h is not None]
            continue
        if not non_null:
            continue
        data.append(row[: len(header)])
    return header or [], data


def convert_value(value, col):
    """Coerce an excel cell to something pymssql can bind. Raises ValueError on bad data."""
    ctype, nullable = col["type"], col["nullable"]

    if isinstance(value, str) and value.strip() == "" and ctype not in TEXTUAL_TYPES:
        value = None

    if value is None:
        if nullable:
            return None
        if ctype in TEXTUAL_TYPES:
            return ""  # legacy char columns are blank-padded, never null
        raise ValueError(f"null in NOT NULL {ctype} column")

    if ctype == "bit":
        if isinstance(value, bool):
            return int(value)
        s = str(value).strip().upper()
        if s in ("1", "1.0", "Y", "YES", "TRUE", "T"):
            return 1
        if s in ("0", "0.0", "N", "NO", "FALSE", "F"):
            return 0
        raise ValueError(f"cannot read {value!r} as bit")

    if ctype in DATETIME_TYPES:
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value
        if isinstance(value, (int, float)):
            return EXCEL_EPOCH + datetime.timedelta(days=float(value))
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y%m%d", "%Y-%m-%d-%H.%M.%S.%f"):
            try:
                return datetime.datetime.strptime(str(value).strip(), fmt)
            except ValueError:
                continue
        raise ValueError(f"cannot read {value!r} as {ctype}")

    if ctype in TEXTUAL_TYPES:
        if isinstance(value, float) and value.is_integer():
            s = str(int(value))
        elif isinstance(value, (datetime.datetime, datetime.date)):
            s = value.isoformat(sep=" ") if isinstance(value, datetime.datetime) else value.isoformat()
        else:
            s = str(value)
        if col["chars"] > 0 and len(s) > col["chars"]:
            raise ValueError(f"{len(s)} chars overflows {ctype}({col['chars']}): {s!r}")
        return s

    if ctype in NUMERIC_TYPES:
        if isinstance(value, str):
            try:
                return float(value) if "." in value else int(value)
            except ValueError:
                raise ValueError(f"cannot read {value!r} as {ctype}")
        return value

    return value


def row_count(conn, table):
    cur = conn.cursor()
    cur.execute(f"SELECT COUNT(*) FROM dbo.[{table}]")
    return cur.fetchone()[0]


def empty_table(conn, table):
    """TRUNCATE where possible, DELETE where an FK blocks it."""
    cur = conn.cursor()
    try:
        cur.execute(f"TRUNCATE TABLE dbo.[{table}]")
        return "truncate"
    except pymssql.Error:
        conn.rollback()
        cur = conn.cursor()
        cur.execute(f"DELETE FROM dbo.[{table}]")
        return "delete"


def prepare(wb, conn, db, table):
    """Validates a sheet against the live schema. Returns (sql, rows, problems)."""
    header, data = read_sheet(wb, table)
    schema = get_schema(conn, db, table)

    problems = []
    if not schema:
        return None, [], [f"table {db}.dbo.{table} does not exist"]

    unknown = [h for h in header if h not in schema]
    if unknown:
        problems.append(f"excel columns absent from DB: {unknown}")
    identity = [c for c, m in schema.items() if m["identity"]]
    if identity:
        problems.append(f"identity columns present, needs IDENTITY_INSERT: {identity}")
    if problems:
        return None, [], problems

    absent = [c for c, m in schema.items() if c not in header and not m["nullable"]]
    if absent:
        problems.append(f"NOT NULL DB columns missing from excel: {absent}")

    rows_out = []
    for i, row in enumerate(data, start=1):
        converted = []
        for h, v in zip(header, row):
            try:
                converted.append(convert_value(v, schema[h]))
            except ValueError as e:
                converted.append(None)
                if len(problems) < 12:
                    problems.append(f"row {i} col {h}: {e}")
        rows_out.append(tuple(converted))

    col_list = ", ".join(f"[{h}]" for h in header)
    return (f"INSERT INTO dbo.[{table}] ({col_list}) VALUES ", len(header)), rows_out, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--commit", action="store_true", help="actually truncate and insert")
    ap.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    ap.add_argument("--only", action="append", help="limit to these tables (repeatable)")
    ap.add_argument("--env-file", default=str(REPO_ROOT / ".env"),
                    help="env file holding MSSQL_HOST/PORT/USER/SA_PASSWORD")
    args = ap.parse_args()

    load_connection(args.env_file)

    workbook = Path(args.workbook)
    if not workbook.is_file():
        sys.exit(f"workbook not found: {workbook}")

    mode = "COMMIT" if args.commit else "DRY RUN"
    print(f"=== {mode} — {USER}@{HOST}:{PORT} — {workbook.name} ===")

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    jobs = [(db, t, m) for db, t, m in JOBS if not args.only or t in args.only]

    total_rows = 0
    errors = 0
    for db, table, mode in jobs:
        conn = connect(db)
        try:
            if table not in wb.sheetnames:
                print(f"  ERROR {db}.dbo.{table}: no sheet named {table}")
                errors += 1
                continue

            before = row_count(conn, table)
            sql, rows, problems = prepare(wb, conn, db, table)

            print(f"  {db}.dbo.{table}: mode={mode}, sheet={len(rows)} rows, table_before={before}")
            for p in problems:
                print(f"    PROBLEM {p}")
            if problems:
                errors += 1
                continue

            total_rows += len(rows)
            if not args.commit:
                continue

            how = empty_table(conn, table) if (mode == "replace" and before) else None
            prefix, ncols = sql
            chunk = rows_per_statement(ncols)
            one_row = "(" + ", ".join(["%s"] * ncols) + ")"
            cur = conn.cursor()
            for start in range(0, len(rows), chunk):
                batch = rows[start:start + chunk]
                stmt = prefix + ", ".join([one_row] * len(batch))
                cur.execute(stmt, tuple(v for row in batch for v in row))
            conn.commit()
            after = row_count(conn, table)
            expected = len(rows) if mode == "replace" else before + len(rows)
            emptied = f"{how}d {before}, " if how else ""
            print(f"    -> {emptied}inserted {len(rows)}, table_after={after}")
            if after != expected:
                print(f"    PROBLEM row count after load ({after}) != expected ({expected})")
                errors += 1
        finally:
            conn.close()

    unloaded = [s for s in wb.sheetnames if s not in {t for _, t, _ in JOBS}]
    if unloaded:
        print(f"--- sheets in workbook, not in this load: {', '.join(unloaded)}")
    print(f"=== done: rows={total_rows} errors={errors} ===")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
